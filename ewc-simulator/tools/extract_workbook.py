#!/usr/bin/env python3
"""Extract EWC_Simulator_Data.xlsx into data.json.

The workbook is canonical; data.json is a derived, committed view of it.
Re-run after every workbook edit:

    python ewc-simulator/tools/extract_workbook.py            # writes data.json
    python ewc-simulator/tools/extract_workbook.py --check    # validate only, no write

Exit status is non-zero if the integrity check fails (substituting actual
placements does not reproduce every club's official CC total).

Requires: openpyxl (pip install openpyxl)
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "EWC_Simulator_Data.xlsx"
OUT = ROOT / "data.json"

SCHEMA_VERSION = 1
REFERENCE_SHEETS = {"Seed Methodology", "Clubs", "CC Reconciliation", "CC Simulation"}

# Canonical names that mark an entry as not belonging to any club. These
# entries still occupy a finishing position in the draw but never score CC points.
PLACEHOLDER_CLUBS = ["(no club)", "(no club tag)", "Mixed", "Unlisted", "Free Agent"]

# Seed kind per method letter, used by the calibration step (never the engine):
# A/B/C seed numbers, D tiers, E no seed.
KIND_BY_METHOD = {"A": "ranked", "B": "ranked", "C": "ranked", "D": "cohort", "E": "flat"}

SIMULATION_DEFAULTS = {"runs": 20000, "rngSeed": 2026, "payingBracket": 24}

# Calibration inputs. Every title uses the TVI on its workbook tab (approved
# 2026-09-19), except titles in LOCKED_TVI, whose approved TVI is final and is
# never replaced by a workbook value.
LOCKED_TVI = {"cs2": {"tvi": 52, "note": "Approved TVI; weights final. Do not recalibrate."}}
CALIBRATION_SEED = {
    "rule": "weight = base ^ (1/T), T = TVI / tviDivisor. base = 1/seed (A/B/C), tier table value (D), 1 (E).",
    "tviDivisor": 50,
    "provisionalTvi": 50,
    "defaultTierTable": {"High": 3.0, "Medium": 1.0, "Low": 0.4},
}

# Official EWC 2026 Club Championship rules that affect ranking.
CC_RULES = {
    "source": "https://esportsworldcup.com/en/rules-and-regulations",
    "bestResultOnly": "An organisation receives only the highest points obtained by one of its participants in a "
                      "competition; lower results are not redistributed.",
    "eligibilityMinTop8": 2,
    "championRequiresTitleWin": True,
    "tieRule": "Rank by points. On equal points, eligible clubs (2+ scoring top-8 finishes) rank above ineligible "
               "ones. A tie for 1st among clubs eligible to win (eligible + at least one title win) is broken by "
               "better single results (Olympic medal logic); any remaining tie shares the rank. A club without a "
               "title win cannot be ranked 1st.",
}

# Club-name standardisation to the Clubs tab Display Name, approved 2026-09-14.
CANONICAL_MERGES = {
    "BetBoom": "BetBoom Team",
    "paiN": "paiN Gaming",
    "ONIC Esports": "ONIC",
    "NASR eSports": "NASR Esports",
    "AlUla Club Esports": "Al-Ula Club",
    "True Rippers": "True Rippers Esports",
    "KINOTROPE Club": "KINOTROPE gaming",
    "EVOS Divine": "EVOS Esports",
    "Elite Esports Europe": "Elite Esports",
}
# HavoK by Vitality scores as its own club (see Fortnite tab note), so it must
# not resolve to Team Vitality in club search.
ALIAS_REMOVALS = {"Team Vitality": ["HavoK by Vitality"]}


def s(v):
    """Cell value as a stripped string, or None."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = str(v).strip()
    return t if t not in ("", "—", "-") else None


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def parse_placement(raw):
    """'5-8' / '5–8' / 3 / '17' -> (label, from, to)."""
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        n = int(raw)
        return str(n), n, n
    t = str(raw).strip().replace("–", "-").replace("—", "-")
    m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", t)
    if m:
        a, b = int(m.group(1)), int(m.group(2))
        return f"{a}-{b}", a, b
    m = re.fullmatch(r"(\d+)", t)
    if m:
        n = int(m.group(1))
        return str(n), n, n
    raise ValueError(f"Unparseable placement: {raw!r}")


def curve_points(curve, place_from):
    """Workbook rule: LOOKUP(first number of placement, from-place column)."""
    pts = 0
    for band in curve:
        if band["from"] <= place_from:
            pts = band["points"]
    return pts


class Report:
    def __init__(self):
        self.errors: list[str] = []
        self.warnings: list[str] = []

    def error(self, msg):
        self.errors.append(msg)

    def warn(self, msg):
        self.warnings.append(msg)


# ---------------------------------------------------------------- title tabs

def find_row(ws, predicate, start=1, end=None):
    end = end or ws.max_row
    for r in range(start, end + 1):
        if predicate(ws.cell(r, 2).value, r):
            return r
    return None


def section_row(ws, number):
    return find_row(ws, lambda v, r: isinstance(v, str) and v.strip().startswith(f"{number} ·"))


def extract_title(ws, rep: Report):
    name = s(ws["B2"].value) or ws.title
    name = re.sub(r"\s+—\s+EWC 2026$", "", name)

    hdr_row = find_row(ws, lambda v, r: s(v) == "FORMAT")
    if hdr_row is None:
        raise ValueError(f"{ws.title}: no FORMAT header")
    labels = {s(ws.cell(hdr_row, c).value): c for c in range(2, ws.max_column + 1) if s(ws.cell(hdr_row, c).value)}
    val = lambda key: s(ws.cell(hdr_row + 1, labels[key]).value)

    field_label = val("FIELD SIZE")
    tvi_label = val("TVI")
    method_label = val("SEED METHOD")
    tvi_m = re.match(r"(\d+)\s*(?:·\s*(.+))?", tvi_label or "")
    method_m = re.match(r"([A-E])\b", method_label or "")
    if not tvi_m:
        rep.error(f"{ws.title}: cannot parse TVI {tvi_label!r}")
    if not method_m:
        rep.error(f"{ws.title}: cannot parse seed method {method_label!r}")
    method = method_m.group(1) if method_m else None

    # ---- block 1: field & results
    sec1, sec2, sec3 = section_row(ws, 1), section_row(ws, 2), section_row(ws, 3)
    field_hdr = find_row(ws, lambda v, r: s(v) in ("Club", "Entry", "Player", "EWC name"), start=sec1)
    cols: dict[str, int] = {}
    for c in range(2, ws.max_column + 1):
        h = s(ws.cell(field_hdr, c).value)
        if h:
            cols[h] = c  # later duplicates (e.g. second 'Club (canonical)') win: rightmost is the join column
    canon_col = cols["Club (canonical)"]
    seed_col = next((c for h, c in cols.items() if re.fullmatch(r"seed( \(.+\))?", h, re.I)), None)
    cohort_col = cols.get("Cohort")
    tier_col = cols.get("Tier (unranked only)") or cols.get("Tier")
    player_col = cols.get("Player") or cols.get("Duo")
    notes_col = cols.get("Notes") or cols.get("Seed basis") or cols.get("Qualification route")
    elig_col = cols.get("CC eligible?")

    entries = []
    r = field_hdr + 1
    while r < (sec2 or ws.max_row + 1):
        b = ws.cell(r, 2).value
        if isinstance(b, str) and re.match(r"TOTAL (POINTS|ELIGIBLE)", b.strip()):
            break
        if b is None:
            r += 1
            continue
        placement = parse_placement(ws.cell(r, cols["Placement"]).value)
        seed_raw = ws.cell(r, seed_col).value if seed_col else None
        seed = None
        if isinstance(seed_raw, (int, float)):
            seed = int(seed_raw) if float(seed_raw).is_integer() else float(seed_raw)
        cohort = s(ws.cell(r, cohort_col).value) if cohort_col else None
        tier = s(ws.cell(r, tier_col).value) if tier_col else None
        official = ws.cell(r, cols["Official CC pts"]).value
        eligible = s(ws.cell(r, elig_col).value) if elig_col else None
        entries.append({
            "row": r,
            # EA FC tabs lead with the player; the entrant label is then the club column.
            "entrant": s(b) if player_col != 2 else s(ws.cell(r, 3).value),
            "player": s(ws.cell(r, player_col).value) if player_col else None,
            "canonical": s(ws.cell(r, canon_col).value),
            "seed": seed,
            "cohort": cohort,
            "tier": tier,
            "placement": placement[0],
            "placeFrom": placement[1],
            "placeTo": placement[2],
            "sheetPoints": ws.cell(r, cols["Points (auto)"]).value,
            "officialPoints": int(official) if isinstance(official, (int, float)) else None,
            "void": (eligible or "").upper() == "VOID",
            "notes": s(ws.cell(r, notes_col).value) if notes_col else None,
        })
        r += 1

    total_row = r
    total_label = s(ws.cell(total_row, 2).value)
    total_value = next((ws.cell(total_row, c).value for c in range(3, ws.max_column + 1)
                        if isinstance(ws.cell(total_row, c).value, (int, float))), None)

    # ---- block 2: points curve ("From place"/"Place" header may sit in column B or C)
    curve = []
    for rr in range(sec2, (sec3 or ws.max_row + 1)):
        heads = [s(ws.cell(rr, c).value) for c in (2, 3)]
        if any(h in ("From place", "Place") for h in heads):
            k = rr + 1
            while isinstance(ws.cell(k, 3).value, (int, float)):
                curve.append({
                    "from": int(ws.cell(k, 3).value),
                    "points": int(ws.cell(k, 4).value),
                    "band": s(ws.cell(k, 5).value),
                })
                k += 1
            break
    if not curve:
        rep.error(f"{ws.title}: points curve not found")
    notes2 = s(ws.cell(sec2 + 1, 2).value)
    points_note = notes2 if notes2 and not notes2.startswith(("From place", "Place")) else None

    # ---- block 3: seed source & rationale, plus any tab-local cohort weights
    sec3_title = s(ws.cell(sec3, 2).value) if sec3 else None
    rationale = []
    tab_weights = None
    if sec3:
        for rr in range(sec3 + 1, ws.max_row + 1):
            v = ws.cell(rr, 2).value
            if isinstance(v, str) and len(v) > 120:
                rationale.append(v.strip())
            if s(ws.cell(rr, 3).value) == "Strength weight":
                tab_weights = {}
                k = rr + 1
                while s(ws.cell(k, 2).value) in ("High", "Medium", "Low"):
                    tab_weights[s(ws.cell(k, 2).value)] = ws.cell(k, 3).value
                    k += 1

    field_m = re.match(r"(\d+)", field_label or "")
    return {
        "sheet": ws.title,
        "name": name,
        "format": val("FORMAT"),
        "fieldSize": int(field_m.group(1)) if field_m else None,
        "fieldSizeLabel": field_label,
        # The tab's TVI is a proposal until approved in data.json calibration.approved.
        "workbookTvi": int(tvi_m.group(1)) if tvi_m else None,
        "tviBand": tvi_m.group(2) if tvi_m else None,
        "seedMethod": method,
        "seedMethodLabel": method_label,
        "seedKind": KIND_BY_METHOD.get(method),
        "seedSource": val("SEED SOURCE"),
        "placementGranularity": val("PLACEMENT GRANULARITY"),
        "description": s(ws["B3"].value),
        "fieldNote": s(ws.cell(sec1 + 1, 2).value),
        "pointsNote": points_note,
        "seedSectionTitle": sec3_title,
        "seedRationale": rationale,
        "tabCohortWeights": tab_weights,
        "sheetTotalLabel": total_label,
        "sheetTotalPoints": total_value,
        "pointsCurve": curve,
        "entries": entries,
    }


# ---------------------------------------------------------------- reference tabs

def extract_seed_methodology(ws):
    methods, proposals = {}, []
    r = find_row(ws, lambda v, r: s(v) == "Method") + 1
    while s(ws.cell(r, 2).value):
        label = s(ws.cell(r, 2).value)
        letter = label[0]
        methods[letter] = {
            "label": label,
            "confidence": s(ws.cell(r, 3).value),
            "what": s(ws.cell(r, 4).value),
            "whenToUse": s(ws.cell(r, 5).value),
            "example": s(ws.cell(r, 6).value),
            "watchOutFor": s(ws.cell(r, 7).value),
            "engineKind": KIND_BY_METHOD[letter],
        }
        r += 1
    r = find_row(ws, lambda v, r: s(v) == "Title") + 1
    while s(ws.cell(r, 2).value):
        proposals.append({
            "title": s(ws.cell(r, 2).value),
            "tvi": ws.cell(r, 3).value,
            "method": s(ws.cell(r, 4).value),
            "source": s(ws.cell(r, 5).value),
            "notes": s(ws.cell(r, 6).value),
        })
        r += 1
    return {"intro": s(ws["B3"].value), "methods": methods, "proposedPerTitle": proposals}


def extract_clubs(ws, rep: Report):
    by_id: dict[str, dict] = {}
    r = find_row(ws, lambda v, r: s(v) == "club_id") + 1
    while s(ws.cell(r, 2).value):
        cid = s(ws.cell(r, 2).value)
        aliases = [a.strip() for a in (s(ws.cell(r, 4).value) or "").split(";") if a.strip()]
        row = {
            "id": cid,
            "name": s(ws.cell(r, 3).value),
            "aliases": aliases,
            "titlesFielded": ws.cell(r, 5).value,
            "qualified": s(ws.cell(r, 6).value),
            "registryCcPoints": ws.cell(r, 7).value,
            "ccTop24": s(ws.cell(r, 8).value) == "Yes",
            "sheetRows": [r],
        }
        if cid in by_id:
            prev = by_id[cid]
            rep.warn(f"Clubs tab: duplicate club_id '{cid}' at rows {prev['sheetRows'] + [r]} "
                     f"(titlesFielded {prev['titlesFielded']} vs {row['titlesFielded']}); merged, kept the larger count")
            prev["aliases"] = sorted(set(prev["aliases"]) | set(aliases))
            prev["titlesFielded"] = max(prev["titlesFielded"] or 0, row["titlesFielded"] or 0)
            prev["sheetRows"].append(r)
        else:
            by_id[cid] = row
        r += 1
    return list(by_id.values())


def extract_reconciliation(ws, title_by_sheet, rep: Report):
    hdr = find_row(ws, lambda v, r: s(v) == "Club")
    # Map each title column to its sheet via the formula in the first data row.
    wsf = openpyxl.load_workbook(WORKBOOK, data_only=False)[ws.title]
    col_title = {}
    for c in range(3, ws.max_column + 1):
        h = s(ws.cell(hdr, c).value)
        if h in ("Sum across tabs", "Official CC total", "Variance", "Status") or not h:
            continue
        formula = str(wsf.cell(hdr + 1, c).value or "")
        m = re.search(r"'([^']+)'!\$", formula) or re.search(r"([A-Za-z0-9 ]+)!\$", formula)
        sheet = m.group(1).strip() if m else None
        if sheet not in title_by_sheet:
            rep.error(f"CC Reconciliation: column {h!r} does not reference a known title sheet ({formula[:60]})")
            continue
        col_title[c] = (h, title_by_sheet[sheet])
    cols = {s(ws.cell(hdr, c).value): c for c in range(2, ws.max_column + 1)}

    rows = []
    r = hdr + 1
    while s(ws.cell(r, 2).value):
        per_title = {}
        for c, (_, tid) in col_title.items():
            v = ws.cell(r, c).value
            per_title[tid] = int(v) if isinstance(v, (int, float)) else 0
        official = ws.cell(r, cols["Official CC total"]).value
        rows.append({
            "club": s(ws.cell(r, 2).value),
            "perTitle": per_title,
            "sheetSum": ws.cell(r, cols["Sum across tabs"]).value,
            "officialTotal": int(official) if isinstance(official, (int, float)) else None,
            "status": s(ws.cell(r, cols["Status"]).value),
        })
        r += 1

    voids = []
    vr = find_row(ws, lambda v, r: isinstance(v, str) and v.startswith("VOIDED CLUB POINTS"), start=r)
    if vr:
        k = find_row(ws, lambda v, r: s(v) == "Club", start=vr) + 1
        while s(ws.cell(k, 2).value):
            voids.append({
                "club": s(ws.cell(k, 2).value),
                "title": s(ws.cell(k, 3).value),
                "footnotePoints": ws.cell(k, 4).value,
                "appliedPoints": ws.cell(k, 5).value,
                "reason": s(ws.cell(k, 6).value),
                "note": s(ws.cell(k, 8).value),
            })
            k += 1
    short_names = {tid: h for c, (h, tid) in col_title.items()}
    order = [tid for c, (h, tid) in sorted(col_title.items())]
    return rows, voids, short_names, order


def extract_reference_run(ws):
    """The workbook's own earlier 20k-run simulation, kept as a comparison baseline."""
    out = {"note": s(ws["B3"].value), "clubs": [], "titles": []}
    r = find_row(ws, lambda v, r: s(v) == "Club") + 1
    while s(ws.cell(r, 2).value):
        out["clubs"].append({
            # Percentages, not fractions: Python writes 5e-05 where JSON.stringify writes 0.00005.
            "club": s(ws.cell(r, 2).value), "pTop8Pct": round(ws.cell(r, 3).value * 100, 6),
            "pWinPct": round(ws.cell(r, 4).value * 100, 6),
            "meanPoints": ws.cell(r, 5).value, "p10": ws.cell(r, 6).value, "p90": ws.cell(r, 7).value,
            "actualPoints": ws.cell(r, 8).value, "actualRank": ws.cell(r, 9).value,
        })
        r += 1
    r = find_row(ws, lambda v, r: s(v) == "Title") + 1
    while s(ws.cell(r, 2).value):
        out["titles"].append({
            "title": s(ws.cell(r, 2).value), "tvi": ws.cell(r, 3).value, "seedType": s(ws.cell(r, 5).value),
            "field": ws.cell(r, 6).value, "effTopBottom": ws.cell(r, 7).value, "topSeedWin": ws.cell(r, 8).value,
        })
        r += 1
    return out


# ---------------------------------------------------------------- validation

def validate(data, rep: Report):
    placeholders = set(data["placeholderClubs"])
    registry_names = {c["name"] for c in data["clubs"]}

    for t in data["titles"]:
        for e in t["entries"]:
            where = f"{t['sheet']} row {e['row']} ({e['entrant']})"
            pts = curve_points(t["pointsCurve"], e["placeFrom"])
            if e["sheetPoints"] != pts:
                rep.error(f"{where}: curve lookup gives {pts}, sheet 'Points (auto)' shows {e['sheetPoints']}")
            scored = 0 if e["void"] else pts
            if e["officialPoints"] is not None and e["officialPoints"] != scored:
                rep.error(f"{where}: engine scores {scored}, 'Official CC pts' is {e['officialPoints']}")
            if not e["canonical"]:
                rep.error(f"{where}: missing canonical club")
            elif e["canonical"] not in placeholders and e["canonical"] not in registry_names:
                rep.error(f"{where}: canonical '{e['canonical']}' is not in the club registry")
            if t["seedKind"] == "ranked" and e["seed"] is None and not e["tier"]:
                rep.error(f"{where}: ranked title but entry has neither numeric seed nor tier")
            if t["seedKind"] == "cohort" and e["cohort"] not in (t["tierTable"] or {}):
                rep.error(f"{where}: cohort title but cohort {e['cohort']!r} is not in the title's tier table")
        if len(t["entries"]) != t["fieldSize"]:
            rep.error(f"{t['sheet']}: {len(t['entries'])} entries but field size {t['fieldSize']}")
        total = sum(0 if e["void"] else curve_points(t["pointsCurve"], e["placeFrom"]) for e in t["entries"])
        if t["sheetTotalPoints"] is not None and total != t["sheetTotalPoints"]:
            rep.error(f"{t['sheet']}: entries score {total}, sheet total row says {t['sheetTotalPoints']}")

    # Integrity check: actual placements through the engine's scoring rules.
    engine_totals = defaultdict(int)
    engine_per_title = defaultdict(dict)
    for t in data["titles"]:
        best = defaultdict(int)
        for e in t["entries"]:
            if e["canonical"] in placeholders:
                continue
            pts = 0 if e["void"] else curve_points(t["pointsCurve"], e["placeFrom"])
            best[e["canonical"]] = max(best[e["canonical"]], pts)
        for club, pts in best.items():
            engine_totals[club] += pts
            engine_per_title[club][t["id"]] = pts

    recon_clubs = set()
    for row in data["reconciliation"]:
        club = row["club"]
        recon_clubs.add(club)
        official = row["officialTotal"] or 0
        if engine_totals.get(club, 0) != official:
            rep.error(f"Integrity: {club} scores {engine_totals.get(club, 0)} from actual placements, official total {official}")
        for tid, v in row["perTitle"].items():
            if engine_per_title[club].get(tid, 0) != v:
                rep.error(f"Integrity: {club} / {tid}: engine {engine_per_title[club].get(tid, 0)}, reconciliation sheet {v}")
    for club, pts in engine_totals.items():
        if pts > 0 and club not in recon_clubs:
            rep.warn(f"{club} scores {pts} from actual placements but is absent from the official CC standings "
                     f"(no reconciliation row); simulated as a normal club")

    # Near-duplicate canonical names: likely the same club split in two, which
    # would split its simulated total. Not fatal (they scored 0), but flag.
    def norm(n):
        n = n.lower().replace("esports", "").replace("e-sports", "").replace("gaming", "").replace("team", "")
        return re.sub(r"[^a-z0-9]", "", n)
    by_norm = defaultdict(set)
    for c in data["clubs"]:
        by_norm[norm(c["name"])].add(c["name"])
    for names in by_norm.values():
        if len(names) > 1:
            rep.warn(f"Possible split club (near-identical canonical names): {sorted(names)}")

    # Seed Methodology proposals vs the title tabs (tabs are authoritative).
    for p in data["seedMethodology"]["proposedPerTitle"]:
        t = next((t for t in data["titles"] if p["title"].lower().startswith(t["name"].lower()[:8])), None)
        if t and (p["method"] != t["seedMethod"] or p["tvi"] != t["workbookTvi"]):
            rep.warn(f"Seed Methodology proposal for {p['title']} ({p['method']}, TVI {p['tvi']}) differs from its tab "
                     f"({t['seedMethod']}, TVI {t['workbookTvi']}); tab value used")
    for tid, a in data["calibration"]["approved"].items():
        t = next((t for t in data["titles"] if t["id"] == tid), None)
        if not t:
            rep.error(f"Calibration approval for unknown title '{tid}'")
        elif a["tvi"] != t["workbookTvi"]:
            rep.warn(f"{tid}: approved TVI {a['tvi']} differs from the workbook tab TVI {t['workbookTvi']}")
    return engine_totals


# ---------------------------------------------------------------- main

def build(rep: Report):
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    raw_titles = [extract_title(ws, rep) for ws in wb.worksheets if ws.title not in REFERENCE_SHEETS]
    sheet_ids = {}
    for t in raw_titles:
        sheet_ids[t["sheet"]] = None
    recon_rows, voids, short_names, order = extract_reconciliation(
        wb["CC Reconciliation"], {t["sheet"]: t["sheet"] for t in raw_titles}, rep)
    # Title ids come from the reconciliation's short column names (CS2, SF6, OW2...).
    sheet_to_id = {sheet: slug(short) for sheet, short in short_names.items()}
    titles = []
    for sheet in order:
        t = next(t for t in raw_titles if t["sheet"] == sheet)
        t = {"id": sheet_to_id[sheet], "short": short_names[sheet], **t}
        for i, e in enumerate(t["entries"]):
            e["id"] = f"{t['id']}:{i}"
            if e["canonical"] in CANONICAL_MERGES:
                e["canonicalSheet"] = e["canonical"]
                e["canonical"] = CANONICAL_MERGES[e["canonical"]]
        # Method D titles carry their own tier table: the tab's strength weights
        # where the tab defines them, otherwise the default 3 / 1 / 0.4.
        tab_weights = t.pop("tabCohortWeights")
        if t["seedKind"] == "cohort":
            t["tierTable"] = {k: float(v) for k, v in tab_weights.items()} if tab_weights \
                else dict(CALIBRATION_SEED["defaultTierTable"])
            t["tierTableSource"] = "title tab" if tab_weights else "default"
        else:
            t["tierTable"], t["tierTableSource"] = None, None
        titles.append(t)
    for row in recon_rows:
        row["perTitle"] = {sheet_to_id[k]: v for k, v in row["perTitle"].items()}
    missing = [t["sheet"] for t in raw_titles if t["sheet"] not in order]
    if missing:
        rep.error(f"Title tabs with no reconciliation column: {missing}")

    # Club registry: Clubs tab first (carries aliases), then reconciliation-only
    # clubs, then any canonical name that appears only in a title tab.
    clubs = extract_clubs(wb["Clubs"], rep)
    names = {c["name"] for c in clubs}
    for row in recon_rows:
        if row["club"] not in names:
            clubs.append({"id": slug(row["club"]).replace("-", "_"), "name": row["club"], "aliases": [],
                          "source": "reconciliation"})
            names.add(row["club"])
    for c in clubs:
        c["aliases"] = [a for a in c["aliases"] if a not in ALIAS_REMOVALS.get(c["name"], [])]
    for t in titles:
        for e in t["entries"]:
            n = e["canonical"]
            if n and n not in names and n not in PLACEHOLDER_CLUBS:
                clubs.append({"id": slug(n).replace("-", "_"), "name": n, "aliases": [], "source": "title-tab"})
                names.add(n)
    for c in clubs:
        c.setdefault("source", "clubs-tab")
        # Entrant names that differ from the canonical name are observed aliases.
        seen = {e["entrant"] for t in titles for e in t["entries"] if e["canonical"] == c["name"]}
        c["entrantNames"] = sorted(n for n in seen if n and n != c["name"])
    seen_ids = set()
    for c in clubs:
        base, n = c["id"], 2
        while c["id"] in seen_ids:
            c["id"] = f"{base}_{n}"
            n += 1
        if c["id"] != base:
            rep.warn(f"Club id '{base}' already used; '{c['name']}' stored as '{c['id']}'")
        seen_ids.add(c["id"])

    # Investment rows are not in the workbook, so they are carried across re-extractions.
    previous = json.loads(OUT.read_text(encoding="utf-8")) if OUT.exists() else {}
    investment = previous.get("investment", [])
    if "Investment" in wb.sheetnames:
        rep.warn("Investment tab found but no reader is implemented yet; keeping data.json investment rows")
    calibration = dict(CALIBRATION_SEED)
    calibration["approved"] = {
        t["id"]: dict(LOCKED_TVI[t["id"]]) if t["id"] in LOCKED_TVI
        else {"tvi": t["workbookTvi"], "note": "Workbook tab TVI."}
        for t in titles
    }

    data = {
        "schemaVersion": SCHEMA_VERSION,
        "source": {
            "workbook": WORKBOOK.name,
            "sha256": hashlib.sha256(WORKBOOK.read_bytes()).hexdigest(),
            "event": "Esports World Cup 2026 — Club Championship",
        },
        "simulationDefaults": SIMULATION_DEFAULTS,
        "calibration": calibration,
        "ccRules": CC_RULES,
        "placeholderClubs": PLACEHOLDER_CLUBS,
        "seedMethodology": extract_seed_methodology(wb["Seed Methodology"]),
        "titles": titles,
        "clubs": clubs,
        "reconciliation": recon_rows,
        "voids": voids,
        "investment": investment,
        "workbookReferenceRun": extract_reference_run(wb["CC Simulation"]),
    }
    return data


def ints_where_whole(v):
    """3.0 -> 3, matching how JavaScript serialises numbers."""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, dict):
        return {k: ints_where_whole(x) for k, x in v.items()}
    if isinstance(v, list):
        return [ints_where_whole(x) for x in v]
    return v


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true", help="validate only; do not write data.json")
    args = ap.parse_args()
    rep = Report()
    data = build(rep)
    validate(data, rep)

    n_entries = sum(len(t["entries"]) for t in data["titles"])
    print(f"titles={len(data['titles'])} entries={n_entries} clubs={len(data['clubs'])} "
          f"reconciliation={len(data['reconciliation'])} voids={len(data['voids'])}")
    for w in rep.warnings:
        print("WARN ", w)
    for e in rep.errors:
        print("ERROR", e)
    if rep.errors:
        print(f"\n{len(rep.errors)} error(s); data.json not written.")
        return 1
    if not args.check:
        # Same bytes as the admin panel's JSON.stringify(data, null, 1) export, so diffs stay clean.
        OUT.write_text(json.dumps(ints_where_whole(data), ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")
        print(f"wrote {OUT.relative_to(ROOT.parent)}")
    print("Integrity check passed: actual placements reproduce every official CC total.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
